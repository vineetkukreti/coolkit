"""Application tracking — SQLite primary, CSV/Excel export optional."""

import csv
import os
from datetime import datetime
from pathlib import Path

from db import init_db, log_application as db_log_application, get_applied_urls as db_get_applied_urls

LOG_FILE = "data/applications_log.csv"
EXCEL_FILE = "data/applications_tracker.xlsx"

HEADERS = [
    "timestamp", "platform", "title", "company", "location",
    "url", "status", "notes", "keywords", "date",
]


def _ensure_log():
    """Create the CSV log file with headers if it doesn't exist."""
    if not os.path.exists(LOG_FILE):
        with open(LOG_FILE, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=HEADERS)
            writer.writeheader()


def log_application(
    platform: str, title: str, company: str,
    location: str, url: str, status: str, notes: str = "",
    keywords: str = "",
):
    """Log a job application to SQLite + CSV backup."""
    # Primary: SQLite
    conn = init_db()
    db_log_application(
        conn, platform, url, title, company, location,
        status=status, apply_method="manual",
    )

    # Backup: CSV
    _ensure_log()
    now = datetime.now()
    row = {
        "timestamp": now.isoformat(),
        "platform": platform,
        "title": title,
        "company": company,
        "location": location,
        "url": url,
        "status": status,
        "notes": notes,
        "keywords": keywords,
        "date": now.strftime("%Y-%m-%d"),
    }
    with open(LOG_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS)
        writer.writerow(row)
    _refresh_excel()


def get_applied_urls() -> set:
    """Return set of URLs we've already applied to (from SQLite)."""
    conn = init_db()
    return db_get_applied_urls(conn)


def _refresh_excel():
    """Refresh the Excel tracker from CSV (optional, fails silently)."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        _ensure_log()
        rows = []
        with open(LOG_FILE, "r", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Applications"
        ws.append(HEADERS)

        header_fill = PatternFill("solid", fgColor="1B4F72")
        header_font = Font(bold=True, color="FFFFFF")
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        status_colors = {
            "applied": "D5F5E3",
            "skipped": "FEF9E7",
            "failed": "FADBD8",
            "partial": "FAE5D3",
        }

        for row in rows:
            values = [row.get(h, "") for h in HEADERS]
            ws.append(values)
            row_num = ws.max_row
            status = row.get("status", "")
            if status in status_colors:
                fill = PatternFill("solid", fgColor=status_colors[status])
                for col in range(1, len(HEADERS) + 1):
                    ws.cell(row=row_num, column=col).fill = fill

        wb.save(EXCEL_FILE)
    except Exception:
        pass
