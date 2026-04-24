#!/usr/bin/env python3
"""Import all lead CSV/XLSX files from leads/ into the startup_leads table."""

import csv
import sqlite3
from pathlib import Path
from db import init_db


LEADS_DIR = Path("data/leads")


def _clean(val):
    """Clean a cell value."""
    if val is None:
        return ""
    return str(val).strip()


def _detect_country(location: str) -> str:
    """Guess country from location string."""
    loc = location.lower()
    if "india" in loc or any(
        state in loc
        for state in [
            "karnataka", "maharashtra", "delhi", "tamil nadu", "telangana",
            "gujarat", "rajasthan", "uttar pradesh", "kerala", "haryana",
            "bengaluru", "mumbai", "hyderabad", "chennai", "pune", "noida",
            "gurugram", "kolkata", "jaipur", "ahmedabad", "lucknow",
        ]
    ):
        return "India"
    if "usa" in loc or "united states" in loc or any(
        s in loc for s in ["california", "new york", "texas", "massachusetts", "florida", "remote"]
    ):
        return "USA"
    if "united kingdom" in loc or "london" in loc or "england" in loc:
        return "UK"
    return ""


def import_healthcare_biotech_192(conn: sqlite3.Connection) -> int:
    """Import healthcare_biotech_startups_192.csv (US startups from YC/etc)."""
    path = LEADS_DIR / "healthcare_biotech_startups_192.csv"
    if not path.exists():
        return 0

    inserted = 0
    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            company = _clean(row.get("Company", ""))
            if not company or company.startswith("Get newly funded"):
                continue

            # Check duplicate
            exists = conn.execute(
                "SELECT 1 FROM startup_leads WHERE company_name = ? AND source_file = ?",
                (company, "healthcare_biotech_startups_192.csv"),
            ).fetchone()
            if exists:
                continue

            desc = _clean(row.get("What They Do", ""))
            location = _clean(row.get("Location", ""))
            website = _clean(row.get("Website", ""))
            email = _clean(row.get("Email(s)", ""))

            # Extract tags from description (AI, Healthcare, Biotech, etc.)
            tags = []
            for tag in ["Artificial Intelligence", "Healthcare", "Biotech", "Medical", "FinTech", "SaaS"]:
                if tag.lower() in desc.lower():
                    tags.append(tag)

            conn.execute(
                """INSERT INTO startup_leads
                   (company_name, website, email, description, industry, location,
                    country, source_file, tags)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    company, website or None, email or None, desc,
                    "Healthcare & Biotech", location,
                    _detect_country(location) or "USA",
                    "healthcare_biotech_startups_192.csv",
                    ", ".join(tags) if tags else None,
                ),
            )
            inserted += 1

    conn.commit()
    return inserted


def import_healthcare_startups_india(conn: sqlite3.Connection) -> int:
    """Import healthcare_startups_india.csv (recently funded Indian startups)."""
    path = LEADS_DIR / "healthcare_startups_india.csv"
    if not path.exists():
        return 0

    inserted = 0
    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            company = _clean(row.get("Company Name", ""))
            if not company:
                continue

            exists = conn.execute(
                "SELECT 1 FROM startup_leads WHERE company_name = ? AND source_file = ?",
                (company, "healthcare_startups_india.csv"),
            ).fetchone()
            if exists:
                continue

            location = _clean(row.get("Location", ""))
            city = location.split(",")[0].strip() if "," in location else location
            desc = _clean(row.get("Description", ""))
            funding = _clean(row.get("Hook", ""))
            email_conf = _clean(row.get("Email Confidence", ""))

            conn.execute(
                """INSERT INTO startup_leads
                   (company_name, description, industry, location, city, country,
                    funding_info, source_file, notes)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    company, desc, "Healthcare",
                    location, city, "India",
                    funding or None,
                    "healthcare_startups_india.csv",
                    f"Email confidence: {email_conf}" if email_conf and email_conf != "Not found" else None,
                ),
            )
            inserted += 1

    conn.commit()
    return inserted


def import_sisfs(conn: sqlite3.Connection) -> int:
    """Import SISFS_Healthcare_Biotech_Startups.csv (DIPP recognized startups)."""
    path = LEADS_DIR / "SISFS_Healthcare_Biotech_Startups.csv"
    if not path.exists():
        return 0

    inserted = 0
    with open(path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            company = _clean(row.get("Company Name", ""))
            if not company:
                continue

            exists = conn.execute(
                "SELECT 1 FROM startup_leads WHERE company_name = ? AND source_file = ?",
                (company, "SISFS_Healthcare_Biotech_Startups.csv"),
            ).fetchone()
            if exists:
                continue

            conn.execute(
                """INSERT INTO startup_leads
                   (company_name, website, industry, sector, city, state, country,
                    employees, funding_status, linkedin_url, facebook_url,
                    twitter_url, youtube_url, other_url, dipp_recognition,
                    entity_type, source_file)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    company,
                    _clean(row.get("Website", "")) or None,
                    _clean(row.get("Industry", "")),
                    _clean(row.get("Sector", "")) or None,
                    _clean(row.get("City", "")) or None,
                    _clean(row.get("State", "")) or None,
                    "India",
                    _clean(row.get("Employees", "")) or None,
                    "Funded" if _clean(row.get("Funding Status", "")) == "Y" else "Not funded",
                    _clean(row.get("LinkedIn", "")) or None,
                    _clean(row.get("Facebook", "")) or None,
                    _clean(row.get("Twitter", "")) or None,
                    _clean(row.get("YouTube", "")) or None,
                    _clean(row.get("Other Link", "")) or None,
                    _clean(row.get("DIPP Recognition", "")) or None,
                    _clean(row.get("Nature of Entity", "")) or None,
                    "SISFS_Healthcare_Biotech_Startups.csv",
                ),
            )
            inserted += 1

    conn.commit()
    return inserted


def import_recruiter_xlsx(conn: sqlite3.Connection, filename: str) -> int:
    """Import recruiter_leads.xlsx (recruiter/hiring posts with emails)."""
    import openpyxl

    path = LEADS_DIR / filename
    if not path.exists():
        return 0

    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active

    headers = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col_map = {h: i for i, h in enumerate(headers) if h}

    def get(row, header, default=""):
        idx = col_map.get(header)
        if idx is not None and idx < len(row) and row[idx] is not None:
            return str(row[idx]).strip()
        return default

    inserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        company = get(row, "Company") or get(row, "Author Name", "")
        if not company:
            continue

        email = get(row, "Email", "")
        if not email:
            continue

        exists = conn.execute(
            "SELECT 1 FROM startup_leads WHERE company_name = ? AND email = ?",
            (company, email),
        ).fetchone()
        if exists:
            continue

        snippet = get(row, "Post Snippet (200 chars)", "")
        role = get(row, "Role Hiring For", "")
        location = get(row, "Location", "")

        conn.execute(
            """INSERT INTO startup_leads
               (company_name, email, description, industry, location, country,
                source_file, tags, notes)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                company, email, snippet,
                "Healthcare & AI",
                location, _detect_country(location),
                filename,
                f"Hiring: {role}" if role else None,
                f"Role: {role}, Experience: {get(row, 'Experience Required', 'N/A')}" if role else None,
            ),
        )
        inserted += 1

    wb.close()
    conn.commit()
    return inserted


def import_all():
    """Import all lead files into the database."""
    conn = init_db()

    print("Importing leads into database...\n")

    n1 = import_healthcare_biotech_192(conn)
    print(f"  healthcare_biotech_startups_192.csv: {n1} leads imported")

    n2 = import_healthcare_startups_india(conn)
    print(f"  healthcare_startups_india.csv:       {n2} leads imported")

    n3 = import_sisfs(conn)
    print(f"  SISFS_Healthcare_Biotech_Startups:   {n3} leads imported")

    n4 = import_recruiter_xlsx(conn, "recruiter_leads.xlsx")
    print(f"  recruiter_leads.xlsx:                {n4} leads imported")

    n5 = import_recruiter_xlsx(conn, "recruiter_leads (1).xlsx")
    print(f"  recruiter_leads (1).xlsx:            {n5} leads imported")

    total = n1 + n2 + n3 + n4 + n5
    print(f"\n  Total: {total} leads imported")

    # Print stats
    stats = conn.execute("SELECT COUNT(*) FROM startup_leads").fetchone()[0]
    print(f"  Total in DB: {stats}")

    by_source = conn.execute(
        "SELECT source_file, COUNT(*) as cnt FROM startup_leads GROUP BY source_file ORDER BY cnt DESC"
    ).fetchall()
    print("\n  By source:")
    for row in by_source:
        print(f"    {row[0]:45s}  {row[1]}")

    by_country = conn.execute(
        "SELECT country, COUNT(*) as cnt FROM startup_leads GROUP BY country ORDER BY cnt DESC"
    ).fetchall()
    print("\n  By country:")
    for row in by_country:
        print(f"    {row[0] or 'Unknown':20s}  {row[1]}")

    conn.close()


if __name__ == "__main__":
    import_all()
